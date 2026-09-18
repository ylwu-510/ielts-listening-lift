#!/usr/bin/env python3
"""Render reviewed IELTS learning content as DOCX/MD without selecting or dropping items."""
import argparse
import importlib.util
import io
import json
import os
from pathlib import Path
import re
import sys
import tempfile

SECTIONS = ("oral", "collocations", "patterns", "passages")
NOTES = {"patterns": "新句式的使用重点", "reading": "阅读中容易误解的表达", "boundaries": "搭配的使用边界"}
EMPTY = "本次材料未另列此类条目。"
BLACK, INK, GRAY, ACCENT = "000000", "182326", "536168", "216A68"


def keys(value, required, optional, path):
    if not isinstance(value, dict):
        raise ValueError(f"{path}: expected an object")
    missing, extra = set(required) - value.keys(), value.keys() - set(required) - set(optional)
    if missing or extra:
        raise ValueError(f"{path}: missing={sorted(missing)}, unknown={sorted(extra)}")


def string(value, path, rich=False):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{path}: expected a nonempty string")
    if any(ord(c) < 32 for c in value):
        raise ValueError(f"{path}: use one paragraph without tabs or control characters")
    if (rich and value.count("**") % 2) or (not rich and "**" in value):
        raise ValueError(f"{path}: invalid bold markers")


def array(value, path):
    if not isinstance(value, list):
        raise ValueError(f"{path}: expected an array")


def validate(data):
    keys(data, ("schema_version", "title", "source_label", "intro", *SECTIONS, "notes"), (), "root")
    if type(data["schema_version"]) is not int or data["schema_version"] != 1:
        raise ValueError("schema_version must be 1")
    for field in ("title", "source_label", "intro"):
        string(data[field], field)
    for group in SECTIONS[:3]:
        array(data[group], group)
        for n, item in enumerate(data[group], 1):
            path = f"{group}[{n}]"
            keys(item, ("phrase", "meaning", "examples"), ("source",), path)
            for field in ("phrase", "meaning"):
                string(item[field], f"{path}.{field}")
            if "source" in item:
                source = item["source"]
                keys(source, ("part", "anchor", "relation"), (), f"{path}.source")
                for field in source:
                    string(source[field], f"{path}.source.{field}")
                if source["relation"] not in ("exact", "adapted"):
                    raise ValueError(f"{path}.source.relation: expected exact or adapted")
            array(item["examples"], f"{path}.examples")
            if not item["examples"]:
                raise ValueError(f"{path}: at least one bilingual example is required")
            for e, example in enumerate(item["examples"], 1):
                keys(example, ("en", "zh"), (), f"{path}.examples[{e}]")
                for field in example:
                    string(example[field], f"{path}.examples[{e}].{field}", rich=True)
    array(data["passages"], "passages")
    for n, item in enumerate(data["passages"], 1):
        keys(item, ("label", "en", "zh"), (), f"passages[{n}]")
        for field in item:
            string(item[field], f"passages[{n}].{field}", rich=field != "label")
    keys(data["notes"], NOTES, (), "notes")
    for group, items in data["notes"].items():
        array(items, f"notes.{group}")
        for n, item in enumerate(items, 1):
            keys(item, ("label", "body"), (), f"notes.{group}[{n}]")
            for field in item:
                string(item[field], f"notes.{group}[{n}].{field}", rich=field == "body")
    return data


def counts(data):
    return {**{s: len(data[s]) for s in SECTIONS}, "notes": {s: len(data["notes"][s]) for s in NOTES}}


def heading(data, section):
    return {"oral": "口语表达", "collocations": "阅读和写作搭配", "patterns": f'{len(data["patterns"])}组句式', "passages": "完整的话"}[section]


def rich_parts(text):
    return [(part, bool(i % 2)) for i, part in enumerate(re.split(r"\*\*(.*?)\*\*", text)) if part]


def md_escape(text):
    return re.sub(r"([\\`*_\[\]<>#])", r"\\\1", text)


def md_rich(text):
    return "".join(f"**{md_escape(s)}**" if bold else md_escape(s) for s, bold in rich_parts(text))


def render_md(data):
    out = [f'# {md_escape(data["title"])}', "", md_escape(data["intro"]), ""]
    for section in SECTIONS:
        out += [f"## {heading(data, section)}", ""]
        if not data[section]:
            out += [EMPTY, ""]
        for n, item in enumerate(data[section], 1):
            prefix = f"{n}. "
            indent = " " * len(prefix)
            if section == "passages":
                out += [f'{prefix}**{md_escape(item["label"])}**  ', f'{indent}{md_rich(item["en"])}  ', f'{indent}{md_rich(item["zh"])}', ""]
            else:
                out += [f'{prefix}**{md_escape(item["phrase"])}**　{md_escape(item["meaning"])}  ']
                for e in item["examples"]:
                    out += [f'{indent}{md_rich(e["en"])}  ', f'{indent}{md_rich(e["zh"])}  ']
                out[-1] = out[-1].rstrip()
                out.append("")
    out += ["## 补充材料", ""]
    for group, title in NOTES.items():
        out += [f"### {title}", ""]
        if not data["notes"][group]:
            out += [EMPTY, ""]
        for n, item in enumerate(data["notes"][group], 1):
            out += [f'{n}. **{md_escape(item["label"])}**　{md_rich(item["body"])}', ""]
    return "\n".join(out).rstrip() + "\n"


def default_cjk():
    return "PingFang SC" if sys.platform == "darwin" else "Microsoft YaHei" if sys.platform == "win32" else "Noto Sans CJK SC"


def render_docx(data, cjk_font=None, latin_font="Arial", paper="letter", breaks=()):
    from docx import Document
    from docx.shared import Inches, Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    cjk_font = cjk_font or default_cjk()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = (Mm(210), Mm(297)) if paper == "a4" else (Inches(8.5), Inches(11))
    sec.top_margin = sec.bottom_margin = Inches(.50)
    sec.left_margin = sec.right_margin = Inches(.62)
    sec.footer_distance = Inches(.19)
    max_count = max([len(data[s]) for s in SECTIONS] + [len(v) for v in data["notes"].values()] + [1])
    text_indent = 19 + max(0, len(str(max_count)) - 2) * 6

    def element(tag, value=None):
        el = OxmlElement("w:" + tag)
        if value is not None:
            el.set(qn("w:val"), str(value))
        return el

    def fonts(rp):
        rf = rp.get_or_add_rFonts()
        for key in list(rf.attrib):
            if "theme" in key.lower():
                del rf.attrib[key]
        for name in ("ascii", "hAnsi", "cs"):
            rf.set(qn("w:" + name), latin_font)
        rf.set(qn("w:eastAsia"), cjk_font)

    specs = [("Normal", 11.5, 14, INK, False), ("Title", 22, 27, BLACK, True),
             ("Heading 1", 16, 21, BLACK, True), ("Heading 2", 12, 17, BLACK, True),
             ("Entry", 10.7, 13.5, BLACK, True), ("Example", 11.5, 14, INK, False),
             ("Translation", 10, 12.5, GRAY, False), ("NotesBody", 10.5, 14, INK, False),
             ("Intro", 10, 13, GRAY, False), ("Passage", 11.5, 15.2, INK, False)]
    for name, size, line, color, bold in specs:
        s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
        if name not in ("Normal", "Title", "Heading 1", "Heading 2"):
            s.base_style = doc.styles["Normal"]
        s.font.name, s.font.size, s.font.bold, s.font.italic = latin_font, Pt(size), bold, False
        s.font.color.rgb = RGBColor.from_string(color)
        fonts(s.element.get_or_add_rPr())
        for border in list(s.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.space_after = Pt(8 if name in ("Title", "Heading 1", "Heading 2") else 0)
        s.paragraph_format.line_spacing = Pt(line)
        s.paragraph_format.widow_control = True
        s.paragraph_format.keep_with_next = name in ("Title", "Heading 1", "Heading 2")

    def run(p, text, bold=None, color=None, size=None):
        r = p.add_run(text)
        fonts(r._element.get_or_add_rPr())
        if bold is not None:
            r.bold = bold
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if size:
            r.font.size = Pt(size)
        return r

    def rich(p, text):
        for part, bold in rich_parts(text):
            run(p, part, True if bold else None, ACCENT if bold else None)

    def paragraph(text, style, after=0, keep=False, indent=None, together=True):
        p = doc.add_paragraph(style=style)
        rich(p, text)
        f = p.paragraph_format
        f.left_indent = Pt(text_indent if indent is None else indent)
        f.space_after, f.keep_with_next, f.keep_together = Pt(after), keep, together
        return p

    numbering = doc.part.numbering_part.element
    aid = max([int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))] + [-1]) + 1
    abstract = element("abstractNum")
    abstract.set(qn("w:abstractNumId"), str(aid))
    abstract.append(element("multiLevelType", "singleLevel"))
    lvl = element("lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1."), ("suff", "tab"), ("lvlJc", "right")):
        lvl.append(element(tag, value))
    pp = element("pPr")
    tabs, tab = element("tabs"), element("tab", "num")
    tab.set(qn("w:pos"), str(text_indent * 20))
    tabs.append(tab)
    pp.append(tabs)
    ind = element("ind")
    ind.set(qn("w:left"), str(text_indent * 20))
    ind.set(qn("w:hanging"), "80")
    pp.append(ind)
    lvl.append(pp)
    rp = element("rPr")
    rf = element("rFonts")
    rf.set(qn("w:ascii"), latin_font)
    rf.set(qn("w:hAnsi"), latin_font)
    rp.extend([rf, element("sz", "21")])
    lvl.append(rp)
    abstract.append(lvl)
    numbering.append(abstract)

    def new_list():
        num = numbering.add_num(aid)
        override = element("lvlOverride")
        override.set(qn("w:ilvl"), "0")
        override.append(element("startOverride", "1"))
        num.append(override)
        return num.get(qn("w:numId"))

    def numbered(style, ident):
        p = doc.add_paragraph(style=style)
        np = element("numPr")
        np.extend([element("ilvl", "0"), element("numId", ident)])
        p._p.get_or_add_pPr().append(np)
        p.paragraph_format.left_indent = Pt(text_indent)
        p.paragraph_format.first_line_indent = Pt(-4)
        return p

    doc.add_paragraph(data["title"], "Title")
    paragraph(data["intro"], "Intro", after=9, indent=0)
    for group in SECTIONS:
        p = doc.add_paragraph(heading(data, group), "Heading 1")
        p.paragraph_format.page_break_before = group in breaks
        p.paragraph_format.space_before = Pt(0 if group == "oral" or group in breaks else 8)
        if not data[group]:
            paragraph(EMPTY, "Translation", after=4, indent=0)
            continue
        ident = new_list()
        for item in data[group]:
            p = numbered("Entry", ident)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_after = Pt(3 if group == "passages" else 1)
            run(p, item["label"] if group == "passages" else item["phrase"], True)
            if group == "passages":
                paragraph(item["en"], "Passage", after=3, together=False)
                paragraph(item["zh"], "Translation", after=8, together=False)
                continue
            if group == "patterns":
                paragraph(item["meaning"], "Translation", after=2, keep=True)
            else:
                run(p, "   " + item["meaning"], False, GRAY, 10)
            for i, e in enumerate(item["examples"]):
                paragraph(e["en"], "Example", keep=True)
                last = i == len(item["examples"]) - 1
                paragraph(e["zh"], "Translation", after=3 if last else 1, keep=not last)
    p = doc.add_paragraph("补充材料", "Heading 1")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.page_break_before = "notes" in breaks
    for group, title in NOTES.items():
        p = doc.add_paragraph(title, "Heading 2")
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
        if not data["notes"][group]:
            paragraph(EMPTY, "Translation", after=4, indent=0)
            continue
        ident = new_list()
        for item in data["notes"][group]:
            p = numbered("NotesBody", ident)
            run(p, item["label"], True)
            run(p, "  ")
            rich(p, item["body"])
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_after = Pt(3)
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(f, data["source_label"] + "   ", color=BLACK, size=8)
    for field in ("PAGE", "NUMPAGES"):
        if field == "NUMPAGES":
            run(f, " / ", color=BLACK, size=8)
        el = element("fldSimple")
        el.set(qn("w:instr"), field)
        f._p.append(el)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = "雅思听力早读学习材料"
    doc.core_properties.author = doc.core_properties.last_modified_by = ""
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def atomic_write(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    name = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as f:
            name = f.name
            f.write(content.encode("utf-8") if isinstance(content, str) else content)
        os.replace(name, path)
    finally:
        if name and os.path.exists(name):
            os.unlink(name)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--format", choices=("auto", "docx", "md", "both"), default="auto")
    parser.add_argument("--cjk-font", default=default_cjk())
    parser.add_argument("--latin-font", default="Arial")
    parser.add_argument("--paper", choices=("letter", "a4"), default="letter")
    parser.add_argument("--break-before", choices=(*SECTIONS[1:], "notes"), action="append", default=[])
    args = parser.parse_args()
    try:
        data = validate(json.loads(args.input.read_text(encoding="utf-8")))
        fmt = args.format
        available = importlib.util.find_spec("docx") is not None
        if fmt == "auto":
            fmt = "docx" if available else "md"
            if not available:
                print("python-docx unavailable; auto selected Markdown with unchanged content.", file=sys.stderr)
        if fmt in ("docx", "both") and not available:
            raise ValueError("DOCX requires python-docx; install requirements.txt or explicitly choose --format md")
        stem = args.output
        if stem.suffix.lower() in (".md", ".docx"):
            stem = stem.with_suffix("")
        outputs = {}
        # Render both before writing either: validation/render errors cannot leave a partial pair.
        if fmt in ("docx", "both"):
            outputs[Path(str(stem) + ".docx")] = render_docx(data, args.cjk_font, args.latin_font, args.paper, args.break_before)
        if fmt in ("md", "both"):
            outputs[Path(str(stem) + ".md")] = render_md(data)
        for path, content in outputs.items():
            atomic_write(path, content)
        print(json.dumps({"format": fmt, "outputs": [str(p.resolve()) for p in outputs], "counts": counts(data)}, ensure_ascii=False))
    except (ValueError, OSError) as error:
        parser.exit(2, f"error: {error}\n")


if __name__ == "__main__":
    main()
